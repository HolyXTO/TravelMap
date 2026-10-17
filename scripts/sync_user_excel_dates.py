import os
import sys
import json
import re
import datetime
import urllib.request
import urllib.error
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

# Configuration
EXCEL_PATH = r"C:\Users\HolyX\.gemini\antigravity\brain\fd8b9b72-acde-47ce-9892-ce301f47fbc2\.user_uploaded\media_1792232919854_cee07191.xlsx"
MAPS_DIR = r"c:\Users\HolyX\Documents\antigravity\travelmap\public\maps"
SUPABASE_URL = "https://yxwflztibppjcyxdulho.supabase.co"
SUPABASE_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6Inl4d2ZsenRpYnBwamN5eGR1bGhvIiwicm9sZSI6ImFub24iLCJpYXQiOjE3ODIzMDYxOTIsImV4cCI6MjA5Nzg4MjE5Mn0.ORs5mJEB5JUNJt6jjHVObSRnoCJtLtXyYabUP5yK9ew"

headers = {
    "apikey": SUPABASE_KEY,
    "Authorization": f"Bearer {SUPABASE_KEY}",
    "Content-Type": "application/json",
    "Prefer": "return=minimal"
}

def supabase_get(path):
    req = urllib.request.Request(f"{SUPABASE_URL}{path}", headers=headers, method="GET")
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode("utf-8"))

import time

def supabase_patch(path, payload, retries=5):
    data_bytes = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    for attempt in range(retries):
        try:
            req = urllib.request.Request(f"{SUPABASE_URL}{path}", headers=headers, data=data_bytes, method="PATCH")
            with urllib.request.urlopen(req, timeout=20) as resp:
                return resp.status
        except Exception as e:
            if attempt == retries - 1:
                raise
            time.sleep(1.5)

def parse_date_cell(val):
    if val is None:
        return []
    if isinstance(val, (datetime.datetime, datetime.date)):
        return [val.strftime("%Y-%m-%d")]
    if isinstance(val, int):
        return [str(val)]
    s = str(val).strip()
    if not s or s == "-":
        return []
    if s in ["居住", "多次"]:
        return [s]
    tokens = [t.strip() for t in re.split(r"[\r\n\s]+", s) if t.strip()]
    dates = []
    for token in tokens:
        if token in ["居住", "多次"]:
            dates.append(token)
            continue
        clean = token.replace("/", "-")
        # Match YYYY-M-D or YYYY-MM-DD
        m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", clean)
        if m:
            dates.append(f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}")
            continue
        # Match YYYY-M or YYYY-MM
        m = re.match(r"^(\d{4})-(\d{1,2})$", clean)
        if m:
            dates.append(f"{m.group(1)}-{int(m.group(2)):02d}")
            continue
        # Match YYYY
        m = re.match(r"^(\d{4})$", clean)
        if m:
            dates.append(m.group(1))
            continue
        # Non-standard string (e.g. 2024-19-18)
        dates.append(clean)

    # Sort standard dates ascending
    calendar_dates = [d for d in dates if d not in ["居住", "多次"]]
    special_tags = [d for d in dates if d in ["居住", "多次"]]
    sorted_dates = sorted(calendar_dates) + special_tags
    return sorted_dates

def main():
    print("⏳ 1. 从 Supabase 拉取当前的 visits 列表...")
    visits_raw = supabase_get("/rest/v1/visits?select=*&order=created_at.desc")
    print(f"   已获取 {len(visits_raw)} 条记录。")

    print("⏳ 2. 加载地图索引与城市数据，对齐 place_groups...")
    with open(os.path.join(MAPS_DIR, "place-index.json"), encoding="utf-8") as f:
        place_index = json.load(f)
    with open(os.path.join(MAPS_DIR, "countries.geojson"), encoding="utf-8") as f:
        countries_geo = json.load(f)
    with open(os.path.join(MAPS_DIR, "states.geojson"), encoding="utf-8") as f:
        states_geo = json.load(f)
    with open(os.path.join(MAPS_DIR, "china-cities.geojson"), encoding="utf-8") as f:
        cities_geo = json.load(f)

    COUNTRY_OVERRIDES = {
        "CHN": "中国", "FRA": "法国", "RUS": "俄罗斯", "SGP": "新加坡",
        "USA": "美国", "DEU": "德国", "GBR": "英国", "ITA": "意大利",
        "ESP": "西班牙", "NOR": "挪威", "SWE": "瑞典", "FIN": "芬兰",
        "NLD": "荷兰", "BEL": "比利时", "CHE": "瑞士", "AUT": "奥地利",
        "GRC": "希腊", "POL": "波兰", "CZE": "捷克", "HUN": "匈牙利"
    }

    place_dict = {}
    for f in countries_geo["features"]:
        p = f["properties"]
        pid = p.get("id") or p.get("code")
        place_dict[pid] = {"localName": COUNTRY_OVERRIDES.get(pid) or p.get("localName") or p.get("name"), "level": "country"}
    for f in states_geo["features"]:
        p = f["properties"]
        place_dict[p.get("id")] = {"localName": p.get("localName") or p.get("name"), "level": "region"}
    for f in cities_geo["features"]:
        p = f["properties"]
        place_dict[p.get("id")] = {"localName": p.get("localName") or p.get("name"), "level": "city"}
    for p in place_index:
        pid = p.get("id")
        if pid:
            existing = place_dict.get(pid, {})
            place_dict[pid] = {**existing, **p}

    SUPPLEMENTAL = [
        {"id": "EXTRA-LONGYEARBYEN", "level": "city", "countryName": "挪威", "localName": "朗伊尔城"},
        {"id": "EXTRA-SVALBARD", "level": "city", "countryName": "挪威", "localName": "斯瓦尔巴群岛"},
        {"id": "EXTRA-SANTORINI", "level": "city", "countryName": "希腊", "localName": "圣托里尼"},
        {"id": "EXTRA-LOFOTEN", "level": "city", "countryName": "挪威", "localName": "罗弗敦群岛"},
        {"id": "EXTRA-BLED", "level": "city", "countryName": "斯洛文尼亚", "localName": "布莱德"},
        {"id": "EXTRA-BUDVA", "level": "city", "countryName": "黑山", "localName": "布德瓦"},
        {"id": "EXTRA-AYIA-NAPA", "level": "city", "countryName": "塞浦路斯", "localName": "圣纳帕"},
        {"id": "EXTRA-VIK", "level": "city", "countryName": "冰岛", "localName": "维克"},
        {"id": "EXTRA-JURMALA", "level": "city", "countryName": "拉脱维亚", "localName": "尤尔马拉"},
        {"id": "EXTRA-MENTON", "level": "city", "countryName": "法国", "localName": "芒通"},
        {"id": "EXTRA-CANNES", "level": "city", "countryName": "法国", "localName": "戛纳"},
        {"id": "EXTRA-RIMINI", "level": "city", "countryName": "意大利", "localName": "里米尼"},
    ]
    for sp in SUPPLEMENTAL:
        place_dict[sp["id"]] = {**place_dict.get(sp["id"], {}), **sp}

    LEGACY_ALIASES = {"it": "ITA", "city-rome": "W-3169070"}

    place_groups = {}
    for v in visits_raw:
        raw_pid = v.get("place_id")
        pid = LEGACY_ALIASES.get(raw_pid, raw_pid)
        if pid not in place_groups:
            local_name = place_dict.get(pid, {}).get("localName") or place_dict.get(pid, {}).get("name") or pid
            place_groups[pid] = {"name_zh": local_name, "visits": []}
        place_groups[pid]["visits"].append(v)

    pids = list(place_groups.keys())
    print(f"   已聚合 {len(pids)} 个地点组。")

    print(f"⏳ 3. 读取用户上传的 Excel 文件：{EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb["地点汇总去重 (154个地点)"]
    print(f"   工作表共 {ws.max_row - 1} 行。")

    updated_count = 0
    skipped_count = 0

    for r in range(2, ws.max_row + 1):
        seq = ws.cell(r, 1).value
        name = ws.cell(r, 2).value
        pid = pids[seq - 1]
        g_visits = place_groups[pid]["visits"]
        
        # Xiao
        x_val = ws.cell(r, 10).value
        x_dates = parse_date_cell(x_val)
        x_rating_val = ws.cell(r, 9).value
        x_comment_val = ws.cell(r, 11).value
        
        # Tang
        t_val = ws.cell(r, 14).value
        t_dates = parse_date_cell(t_val)
        t_rating_val = ws.cell(r, 13).value
        t_comment_val = ws.cell(r, 15).value

        for prof_id, dates, r_val, c_val, p_label in [
            ("person-a", x_dates, x_rating_val, x_comment_val, "Xiao"),
            ("person-b", t_dates, t_rating_val, t_comment_val, "Tang")
        ]:
            v = next((item for item in g_visits if item["profile_id"] == prof_id), None)
            if not v:
                continue
            
            # Old note parse
            old_note = {}
            if v.get("note"):
                try:
                    old_note = json.loads(v["note"])
                    if not isinstance(old_note, dict):
                        old_note = {"text": str(v["note"])}
                except:
                    old_note = {"text": str(v["note"])}
            
            # Rating
            rating = old_note.get("rating", 0)
            if r_val is not None and r_val != "-":
                try:
                    rating = float(r_val)
                except:
                    pass
            
            # Comment
            comment = old_note.get("comment", "")
            if c_val is not None and c_val != "-":
                comment = str(c_val).strip()
                
            text = old_note.get("text", "")
            transport_mode = old_note.get("transportMode", "")
            
            final_dates = [str(d).strip() for d in (dates if dates else old_note.get("visitDates", [])) if d]
            date_display = str(final_dates[-1] if final_dates else (old_note.get("dateDisplay") or "")).strip()
            
            # Precision & db_date
            db_date = None
            precision = "day"
            if date_display == "居住":
                precision = "living"
            elif date_display == "多次":
                precision = "multiple"
            elif not date_display:
                precision = "none"
            elif re.match(r"^\d{4}-\d{2}-\d{2}$", date_display):
                precision = "day"
                db_date = date_display
            elif re.match(r"^\d{4}-\d{2}$", date_display):
                precision = "month"
                db_date = f"{date_display}-01"
            elif re.match(r"^\d{4}$", date_display):
                precision = "year"
                db_date = f"{date_display}-01-01"
            else:
                for d in reversed(final_dates):
                    d_str = str(d).strip()
                    if re.match(r"^\d{4}-\d{2}-\d{2}$", d_str):
                        db_date = d_str
                        break
                    elif re.match(r"^\d{4}-\d{2}$", d_str):
                        db_date = f"{d_str}-01"
                        break
                    elif re.match(r"^\d{4}$", d_str):
                        db_date = f"{d_str}-01-01"
                        break
                precision = "custom"
                
            new_note_payload = {
                "dateDisplay": date_display,
                "datePrecision": precision,
                "rating": rating,
                "text": text,
                "comment": comment,
                "transportMode": transport_mode,
                "visitDates": final_dates
            }

            patch_data = {
                "visited_at": db_date,
                "note": json.dumps(new_note_payload, ensure_ascii=False)
            }

            # Update Supabase
            try:
                supabase_patch(f"/rest/v1/visits?id=eq.{v['id']}", patch_data)
                updated_count += 1
                if final_dates and (len(final_dates) > 1 or final_dates[0] in ["居住", "多次"]):
                    print(f"  ✓ [{name}] {p_label}: 到访时间已更新为 {final_dates} (db_date={db_date})")
            except Exception as e:
                print(f"  ✗ 更新失败 [{name}] {p_label} (ID: {v['id']}): {e}")

    print(f"\n🎉 同步完成！共更新 {updated_count} 条足迹记录到 Supabase！")

if __name__ == "__main__":
    main()
